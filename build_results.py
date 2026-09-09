import json,re,glob,os,sys
from pathlib import Path
from openpyxl import load_workbook
from bijoy2unicode import bijoy2unicode

ROOT=Path(__file__).resolve().parent
EXCEL_DIR=ROOT/'excel'
DATA_DIR=ROOT/'data'
DATA_DIR.mkdir(exist_ok=True)

CLASS_BN={'Class-1':'প্রথম শ্রেণি','Class-2':'দ্বিতীয় শ্রেণি','Class-3':'তৃতীয় শ্রেণি','Class-4':'চতুর্থ শ্রেণি','Class-5':'পঞ্চম শ্রেণি','Class-6':'ষষ্ঠ শ্রেণি','Narsari':'নার্সারি','Hifz':'হিফজ'}
EXAM_INFO={
 '1st-term':('First Term Exam','প্রথম সাময়িক পরীক্ষা'),
 '2nd-term':('Second Term Exam','দ্বিতীয় সাময়িক পরীক্ষা'),
 'annual':('Annual Exam','বার্ষিক পরীক্ষা'),
}

def clean(v):
    if v is None: return ''
    if isinstance(v,str):
        v=v.replace('\xa0',' ').strip()
        if not v: return ''
        # The supplied workbook uses legacy Bijoy/SutonnyMJ text for names and
        # headers, so convert plain ASCII-looking Bengali text as well.
        if re.search(r'[\u0980-\u09ff]',v): return v
        return bijoy2unicode(v).strip()
    return v

def num(v):
    if v is None or v=='': return None
    if isinstance(v,(int,float)) and not isinstance(v,bool): return v
    s=str(v).strip().replace(',','')
    if s in ('*','-','—'): return None
    try:
        x=float(s)
        return int(x) if x.is_integer() else x
    except: return None

def rank_value(v):
    if v is None or v=='': return None
    n=num(v)
    if n is not None: return n
    return clean(v)

def detect_exam_year(path, wb):
    stem=path.stem.lower()
    m=re.search(r'(20\d{2})[-_ ]?(1st|2nd|annual)',stem)
    if m:
        year=m.group(1)
        key={'1st':'1st-term','2nd':'2nd-term','annual':'annual'}[m.group(2)]
        return year,key
    key=None
    if 'first' in stem or '1st' in stem: key='1st-term'
    elif 'second' in stem or '2nd' in stem: key='2nd-term'
    elif 'annual' in stem: key='annual'
    # Fallback to title text in first class sheet.
    for ws in wb.worksheets:
        if ws.title=='Formula_Test': continue
        for row in range(1,min(ws.max_row,14)+1):
            for col in range(1,min(ws.max_column,12)+1):
                v=ws.cell(row,col).value
                if isinstance(v,str) and v.strip():
                    t=clean(v)
                    ym=re.search(r'(20\d{2})',t)
                    if not ym:
                        continue
                    year=ym.group(1)
                    if key is None:
                        if 'দ্বিতীয়' in t or 'দ্বিতীয়' in t: key='2nd-term'
                        elif 'প্রথম' in t: key='1st-term'
                        elif 'বার্ষিক' in t: key='annual'
                    if key: return year,key
    raise ValueError(f'Could not determine year/exam from {path.name}. Use a name like 2027-1st-term.xlsx')

def parse_sheet(ws, year, key):
    if ws.title not in CLASS_BN: return []
    # Header row is the row containing the total-mark header.
    header_row=None; total_col=None
    for r in range(1,min(ws.max_row,30)+1):
        for c in range(1,ws.max_column+1):
            t=clean(ws.cell(r,c).value)
            if isinstance(t,str) and 'সর্ব মোট' in t:
                header_row=r; total_col=c; break
        if header_row: break
    if not header_row or not total_col or total_col<3:
        raise ValueError(f'{ws.title}: could not find the marks header row/total column')
    headers=[clean(ws.cell(header_row,c).value) for c in range(3,total_col)]
    subjects=[h.replace('\n',' ').strip() for h in headers if h]
    # In case of merged/blank header cells, keep positional columns but ignore empty ones.
    col_subject=[]
    for c in range(3,total_col):
        h=str(clean(ws.cell(header_row,c).value) or '').replace('\n',' ').strip()
        if h: col_subject.append((c,h))
    # Metadata columns after total: average, point, grade, position/rank.
    after=[clean(ws.cell(header_row,c).value).replace('\n',' ').strip().lower() for c in range(total_col+1,ws.max_column+1)]
    avg_col=point_col=grade_col=rank_col=None
    for c in range(total_col+1,ws.max_column+1):
        h=str(clean(ws.cell(header_row,c).value) or '').replace('\n',' ').strip()
        hl=h.lower()
        if not h: continue
        if h in ('গড়','গড়') or 'average' in hl: avg_col=c
        elif 'পয়েন্ট' in h or 'পয়েন্ট' in h or 'point' in hl: point_col=c
        elif h in ('গ্রেড','grade') or 'grade' in hl: grade_col=c
        elif 'অবস্থান' in h or 'position' in hl or 'rank' in hl: rank_col=c
    # Usually the columns are fixed after total; use positional fallback.
    if avg_col is None and total_col+1<=ws.max_column: avg_col=total_col+1
    if point_col is None and total_col+2<=ws.max_column: point_col=total_col+2
    if grade_col is None and total_col+3<=ws.max_column: grade_col=total_col+3
    if rank_col is None and total_col+4<=ws.max_column: rank_col=total_col+4

    exam,exam_bn=EXAM_INFO[key]
    out=[]
    for r in range(header_row+1,ws.max_row+1):
        roll=ws.cell(r,1).value
        name=clean(ws.cell(r,2).value)
        if roll is None and not name: continue
        if not name: continue
        # Ignore footer/legend rows that don't have a usable roll.
        roll_num=num(roll)
        if roll_num is None: continue
        subjects_data=[]
        for c,h in col_subject:
            subjects_data.append({'name':h,'marks':num(ws.cell(r,c).value)})
        total=num(ws.cell(r,total_col).value)
        average=num(ws.cell(r,avg_col).value if avg_col else None)
        point=num(ws.cell(r,point_col).value if point_col else None)
        grade=str(ws.cell(r,grade_col).value).strip() if grade_col and ws.cell(r,grade_col).value is not None else ''
        rank=rank_value(ws.cell(r,rank_col).value if rank_col else None)
        out.append({'year':str(year),'exam':exam,'examBn':exam_bn,'className':ws.title,'classBn':CLASS_BN[ws.title],'roll':str(int(roll_num) if isinstance(roll_num,(int,float)) and float(roll_num).is_integer() else roll_num),'reg':'','name':name,'subjects':subjects_data,'total':total,'average':average,'point':point,'grade':grade,'rank':rank})
    return out

def build():
    # Keep the current website data as the initial seed. Any Excel file placed in
    # /excel is then added; if the same year+exam already exists, that exam is replaced.
    result_path=DATA_DIR/'results.json'
    if result_path.exists():
        try:
            base=json.loads(result_path.read_text(encoding='utf-8'))
            all_rows=base if isinstance(base,list) else base.get('students',[])
        except Exception:
            all_rows=[]
    else:
        all_rows=[]
    sources=[]
    files=sorted(EXCEL_DIR.glob('*.xlsx'))+sorted(EXCEL_DIR.glob('*.xlsm'))
    for path in files:
        wb=load_workbook(path,data_only=True,read_only=False)
        year,key=detect_exam_year(path,wb)
        rows=[]
        for ws in wb.worksheets:
            rows.extend(parse_sheet(ws,year,key))
        if not rows:
            raise ValueError(f'{path.name}: no student records were found')
        # Remove the old copy of this year/exam, then insert the fresh Excel data.
        exam_name=EXAM_INFO[key][0]
        all_rows=[s for s in all_rows if not (str(s.get('year'))==str(year) and s.get('exam')==exam_name)]
        all_rows.extend(rows)
        sources.append({'file':path.name,'year':year,'examKey':key,'records':len(rows)})
    all_rows.sort(key=lambda s:(-int(s.get('year',0)), s.get('exam',''), s.get('className',''), int(s.get('roll')) if str(s.get('roll','')).isdigit() else 999999))
    payload={'generatedAt':'AUTO','sources':sources,'students':all_rows}
    result_path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    years=sorted({str(s.get('year')) for s in all_rows},key=int,reverse=True)
    manifest={'years':years,'sources':sources,'classes':list(CLASS_BN.keys())}
    (DATA_DIR/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Generated {len(all_rows)} records total; processed {len(sources)} Excel file(s).')
    for x in sources: print(x)
    return all_rows

if __name__=='__main__': build()
